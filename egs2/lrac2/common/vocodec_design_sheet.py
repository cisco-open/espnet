#!/usr/bin/env python3
# Copyright 2026 Cisco Systems, Inc. and its affiliates
# Apache-2.0

"""Build and verify an LRAC 2 VoCodec model-design workbook."""

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SAMPLE_RATE = 24000
FFT_SIZE = 720
CODEBOOK_DIM = 8
CODEBOOK_SIZE = 1024
BYTES_PER_VALUE = 2
STATE_READS_PER_CALL = 1
STATE_WRITES_PER_CALL = 1

TRANSMIT_BUDGET_MFLOPS = 400
RECEIVE_BUDGET_MFLOPS = 300
LATENCY_BUDGET_MS = 30
MEMORY_BANDWIDTH_PMAX = 15
SUPPORTED_BITRATES_KBPS = (1, 6, 12, 24, 32)

NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
GREEN = "E2F0D9"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7B7B7")


@dataclass(frozen=True)
class Variant:
    track: int
    label: str
    hop_size: int
    attention_context_frames: int
    attention_lookahead_frames: int
    attention_heads: int
    rvq_layers: int
    model_dim: int
    encoder_convnext_expansion_dim: int
    decoder_convnext_expansion_dim: int
    decoder_posnet_groups_first: tuple
    decoder_posnet_groups_last: tuple
    attention_use_rope: bool = False
    attention_rotary_dim: int | None = None
    supported_bitrates_kbps: tuple = SUPPORTED_BITRATES_KBPS
    parameter_budget: int | None = None

    @property
    def frame_rate(self):
        return SAMPLE_RATE / self.hop_size

    def sheet_name(self, component):
        return f"{component}-Track{self.track}"


def style_title(sheet, title, last_column):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = sheet.cell(1, 1, title)
    cell.font = Font(size=16, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 25


def style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(bottom=THIN_GRAY)


def style_table(sheet, first_row, last_row, first_column, last_column):
    for row in sheet.iter_rows(
        min_row=first_row,
        max_row=last_row,
        min_col=first_column,
        max_col=last_column,
    ):
        for cell in row:
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
            )
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(sheet, widths):
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width


@dataclass(frozen=True)
class LayerSpec:
    """One explicitly accounted operation in the encoder or decoder."""

    name: str
    block: str
    instances: int
    operation: str
    kernel_or_context: int
    dilation: int
    input_channels: int
    output_channels: int
    groups_or_heads: int = 1
    mac_kind: str = "none"
    parameter_kind: str = "none"
    state_kind: str = "none"
    persistent_override: int | None = None
    parameter_extra: int = 0
    parameter_extra_note: str = ""
    persistent_extra: int = 0
    left_context: int = 0
    right_context: int = 0
    context_domain: str = "frame"

    def macs_per_position(self):
        if self.mac_kind == "conv":
            return (
                self.kernel_or_context
                * self.input_channels
                * self.output_channels
                // self.groups_or_heads
            )
        if self.mac_kind == "attention":
            # heads * keys * head_dim = keys * channels
            return self.kernel_or_context * self.input_channels
        if self.mac_kind == "rope_phase":
            return self.kernel_or_context
        if self.mac_kind == "rope_rotation":
            # Two multiply-accumulates per rotated scalar, across all heads.
            return 2 * self.kernel_or_context * self.groups_or_heads
        return 0

    def trainable_parameters(self):
        if self.parameter_kind == "conv":
            return (
                self.macs_per_position() + self.output_channels + self.parameter_extra
            )
        if self.parameter_kind == "norm":
            return 2 * self.output_channels + self.parameter_extra
        return self.parameter_extra

    def persistent_values(self, variant):
        if self.persistent_override is not None:
            return self.persistent_override + self.persistent_extra
        return self.trainable_parameters() + self.persistent_extra

    def state_values(self, variant):
        if self.state_kind == "causal_conv":
            effective_kernel = (self.kernel_or_context - 1) * self.dilation + 1
            state_size = (effective_kernel - 1) * self.input_channels
            return (STATE_READS_PER_CALL + STATE_WRITES_PER_CALL) * state_size
        if self.state_kind in ("attention_k", "attention_v"):
            # Context includes the current frame. Read the cached past/lookahead
            # tensor, then write the current tensor.
            cache_reads = (
                STATE_READS_PER_CALL
                * (
                    variant.attention_context_frames
                    - 1
                    + variant.attention_lookahead_frames
                )
                * self.input_channels
            )
            cache_writes = STATE_WRITES_PER_CALL * self.input_channels
            return cache_reads + cache_writes
        if self.state_kind == "stft":
            state_size = FFT_SIZE - variant.hop_size
            return (STATE_READS_PER_CALL + STATE_WRITES_PER_CALL) * state_size
        if self.state_kind == "rope_inv_freq":
            # Read the derived inverse-frequency vector once per invocation.
            # Generated sine/cosine values are retained for both Q and K.
            return STATE_READS_PER_CALL * (variant.attention_rotary_dim // 2)
        return 0


def conv_layer(
    name,
    block,
    instances,
    input_channels,
    output_channels,
    kernel_size=1,
    groups=1,
    causal=False,
    parameter_extra=0,
    parameter_extra_note="",
    persistent_extra=0,
):
    return LayerSpec(
        name=name,
        block=block,
        instances=instances,
        operation="Causal Conv1d" if causal else "Conv1d / linear projection",
        kernel_or_context=kernel_size,
        dilation=1,
        input_channels=input_channels,
        output_channels=output_channels,
        groups_or_heads=groups,
        mac_kind="conv",
        parameter_kind="conv",
        state_kind="causal_conv" if causal and kernel_size > 1 else "none",
        parameter_extra=parameter_extra,
        parameter_extra_note=parameter_extra_note,
        persistent_extra=persistent_extra,
        left_context=kernel_size - 1 if causal else 0,
    )


def convnext_layers(hidden_channels, expansion_channels):
    block = "ConvNeXt blocks 1–6"
    instances = 6
    return [
        conv_layer(
            "Depthwise convolution",
            block,
            instances,
            hidden_channels,
            hidden_channels,
            kernel_size=7,
            groups=hidden_channels,
            causal=True,
            parameter_extra=2 * hidden_channels,
            parameter_extra_note="ConvNeXt LayerNorm scale and bias",
        ),
        conv_layer(
            "Pointwise input projection",
            block,
            instances,
            hidden_channels,
            expansion_channels,
        ),
        conv_layer(
            "Pointwise output projection",
            block,
            instances,
            expansion_channels,
            hidden_channels,
        ),
    ]


def posnet_resnet_layers(block, instances, hidden_dim, groups):
    return [
        conv_layer(
            "Residual convolution 1",
            block,
            instances,
            hidden_dim,
            hidden_dim,
            kernel_size=3,
            groups=groups[0],
            causal=True,
            parameter_extra=2 * hidden_dim,
            parameter_extra_note="PosNet residual-block norm1 scale and bias",
        ),
        conv_layer(
            "Residual convolution 2",
            block,
            instances,
            hidden_dim,
            hidden_dim,
            kernel_size=3,
            groups=groups[1],
            causal=True,
            parameter_extra=2 * hidden_dim,
            parameter_extra_note="PosNet residual-block norm2 scale and bias",
        ),
    ]


def attention_layers(variant):
    block = "Bounded attention after PosNet block 2"
    keys_per_query = (
        variant.attention_context_frames + variant.attention_lookahead_frames
    )
    hidden_dim = variant.model_dim
    layers = [
        conv_layer(
            "QKV projection",
            block,
            1,
            hidden_dim,
            3 * hidden_dim,
            parameter_extra=2 * hidden_dim,
            parameter_extra_note="Attention input LayerNorm scale and bias",
        ),
    ]
    if variant.attention_use_rope:
        layers += [
            LayerSpec(
                name="RoPE phase generation",
                block=block,
                instances=1,
                operation=(
                    "Position × derived inverse frequencies; sine/cosine "
                    "arithmetic excluded"
                ),
                kernel_or_context=variant.attention_rotary_dim // 2,
                dilation=1,
                input_channels=hidden_dim,
                output_channels=hidden_dim,
                mac_kind="rope_phase",
                state_kind="rope_inv_freq",
            ),
            LayerSpec(
                name="RoPE query rotation",
                block=block,
                instances=1,
                operation="Pairwise rotary transform of Q",
                kernel_or_context=variant.attention_rotary_dim,
                dilation=1,
                input_channels=hidden_dim,
                output_channels=hidden_dim,
                groups_or_heads=variant.attention_heads,
                mac_kind="rope_rotation",
            ),
            LayerSpec(
                name="RoPE key rotation",
                block=block,
                instances=1,
                operation="Pairwise rotary transform of K before cache use",
                kernel_or_context=variant.attention_rotary_dim,
                dilation=1,
                input_channels=hidden_dim,
                output_channels=hidden_dim,
                groups_or_heads=variant.attention_heads,
                mac_kind="rope_rotation",
            ),
        ]
    layers += [
        LayerSpec(
            name="Query-key score product",
            block=block,
            instances=1,
            operation="Q × Kᵀ (configured bounded context)",
            kernel_or_context=keys_per_query,
            dilation=1,
            input_channels=hidden_dim,
            output_channels=hidden_dim,
            groups_or_heads=variant.attention_heads,
            mac_kind="attention",
            state_kind="attention_k",
            left_context=variant.attention_context_frames - 1,
            right_context=variant.attention_lookahead_frames,
        ),
        LayerSpec(
            name="Attention-weighted value product",
            block=block,
            instances=1,
            operation="softmax(QKᵀ) × V",
            kernel_or_context=keys_per_query,
            dilation=1,
            input_channels=hidden_dim,
            output_channels=hidden_dim,
            groups_or_heads=variant.attention_heads,
            mac_kind="attention",
            state_kind="attention_v",
        ),
        conv_layer("Attention output projection", block, 1, hidden_dim, hidden_dim),
    ]
    return layers


def stft_state_layer(component):
    """Account for the persistent window and streaming overlap state."""

    return LayerSpec(
        name=f"{component} window and overlap state",
        block=component,
        instances=1,
        operation="Persistent window and streaming state; FFT MACs excluded",
        kernel_or_context=FFT_SIZE,
        dilation=1,
        input_channels=1,
        output_channels=1,
        state_kind="stft",
        persistent_override=FFT_SIZE,
        context_domain="sample",
    )


def encoder_components(variant):
    hidden_dim = variant.model_dim
    layers = [
        stft_state_layer("STFT analysis"),
        conv_layer(
            "STFT feature projection",
            "Encoder input",
            1,
            2 * (FFT_SIZE // 2 + 1),
            hidden_dim,
            parameter_extra=2 * hidden_dim,
            parameter_extra_note="Encoder input LayerNorm scale and bias",
        ),
    ]
    layers += convnext_layers(hidden_dim, variant.encoder_convnext_expansion_dim)
    layers += posnet_resnet_layers("PosNet ResNet blocks 1–2", 2, hidden_dim, (1, 1))
    layers += attention_layers(variant)
    layers += posnet_resnet_layers("PosNet ResNet blocks 3–4", 2, hidden_dim, (1, 1))
    layers += [
        conv_layer(
            "Latent projection",
            "Encoder output",
            1,
            hidden_dim,
            hidden_dim,
            kernel_size=3,
            causal=True,
            parameter_extra=2 * hidden_dim,
            parameter_extra_note="PosNet output LayerNorm scale and bias",
        ),
    ]
    return tuple(layers)


def decoder_components(variant):
    hidden_dim = variant.model_dim
    layers = [
        stft_state_layer("ISTFT synthesis"),
        conv_layer(
            "Latent projection",
            "Decoder input",
            1,
            hidden_dim,
            hidden_dim,
            kernel_size=3,
            causal=True,
        ),
    ]
    layers += posnet_resnet_layers(
        "PosNet ResNet blocks 1–2",
        2,
        hidden_dim,
        variant.decoder_posnet_groups_first,
    )
    layers += attention_layers(variant)
    layers += posnet_resnet_layers(
        "PosNet ResNet blocks 3–4",
        2,
        hidden_dim,
        variant.decoder_posnet_groups_last,
    )
    layers += convnext_layers(hidden_dim, variant.decoder_convnext_expansion_dim)
    layers += [
        conv_layer(
            "Complex-spectrum output projection",
            "Decoder output",
            1,
            hidden_dim,
            2 * (FFT_SIZE // 2 + 1),
            parameter_extra=4 * hidden_dim,
            parameter_extra_note=(
                "PosNet output LayerNorm and decoder output LayerNorm scales and biases"
            ),
        ),
    ]
    return tuple(layers)


def add_component_sheet(
    workbook,
    variant,
    name,
    components,
):
    sheet = workbook.create_sheet(name)
    component_name = name.split("-", 1)[0]
    encoder_sheet = variant.sheet_name("Encoder")
    style_title(sheet, f"VoCodec {component_name} — {variant.label}", 22)
    is_encoder = component_name == "Encoder"
    sampling_rate = SAMPLE_RATE if is_encoder else f"='{encoder_sheet}'!$B$2"
    fft_size = FFT_SIZE if is_encoder else f"='{encoder_sheet}'!$B$3"
    hop_size = variant.hop_size if is_encoder else f"='{encoder_sheet}'!$B$4"
    model_channels = variant.model_dim if is_encoder else f"='{encoder_sheet}'!$B$6"
    expansion_channels = (
        variant.encoder_convnext_expansion_dim
        if is_encoder
        else variant.decoder_convnext_expansion_dim
    )
    qkv_channels = "=3*B6"
    stft_channels = "=2*(B3/2+1)" if is_encoder else f"='{encoder_sheet}'!$B$9"
    attention_context = (
        variant.attention_context_frames if is_encoder else f"='{encoder_sheet}'!$B$10"
    )
    attention_lookahead = (
        variant.attention_lookahead_frames
        if is_encoder
        else f"='{encoder_sheet}'!$B$11"
    )
    attention_heads = (
        variant.attention_heads if is_encoder else f"='{encoder_sheet}'!$B$12"
    )
    attention_use_rope = (
        variant.attention_use_rope if is_encoder else f"='{encoder_sheet}'!$B$13"
    )
    attention_rotary_dim = (
        variant.attention_rotary_dim if is_encoder else f"='{encoder_sheet}'!$B$14"
    )
    bytes_per_value = BYTES_PER_VALUE if is_encoder else f"='{encoder_sheet}'!$B$16"
    state_reads_per_call = (
        STATE_READS_PER_CALL if is_encoder else f"='{encoder_sheet}'!$B$17"
    )
    state_writes_per_call = (
        STATE_WRITES_PER_CALL if is_encoder else f"='{encoder_sheet}'!$B$18"
    )
    transform_latency = (
        (
            "STFT buffering latency (ms)",
            "=$B$4/$B$2*1000",
        )
        if is_encoder
        else (
            "ISTFT algorithmic latency (ms)",
            "=($B$3-$B$4)/$B$2*1000",
        )
    )
    settings = (
        ("Sampling rate (Hz)", sampling_rate),
        ("STFT window / FFT size (samples)", fft_size),
        ("STFT hop (samples)", hop_size),
        ("Overall Stride (sampling rate / STFT hop)", "=B2/B4"),
        ("Model channels", model_channels),
        (
            f"{component_name} ConvNeXt expansion channels",
            expansion_channels,
        ),
        ("Attention QKV channels", qkv_channels),
        ("Complex STFT feature channels", stft_channels),
        (
            "Attention context (frames, current included)",
            attention_context,
        ),
        ("Attention lookahead (frames)", attention_lookahead),
        ("Attention heads", attention_heads),
        ("Attention uses RoPE", attention_use_rope),
        ("RoPE rotary dimensions / head", attention_rotary_dim or 0),
        transform_latency,
        ("Bytes per parameter/buffer value", bytes_per_value),
        ("Runtime state reads per call", state_reads_per_call),
        ("Runtime state writes per call", state_writes_per_call),
    )
    for row, (label, value) in enumerate(settings, 2):
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, value)
        sheet.cell(row, 1).font = Font(bold=True)
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor=BLUE)
        sheet.cell(row, 2).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        sheet.cell(row, 1).border = sheet.cell(row, 2).border = Border(
            left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
        )
    settings_end_row = len(settings) + 1
    sheet.merge_cells(
        start_row=2, start_column=4, end_row=settings_end_row, end_column=22
    )
    sheet.cell(
        2,
        4,
        "Computations only include MAC-bearing operations. The first STFT/ISTFT row has "
        "zero counted MACs but explicitly accounts for its persistent window and "
        "streaming overlap state. LayerNorm arithmetic is omitted, while each "
        "LayerNorm's scale and bias values are included in column L through the "
        "+2*$B$6 terms; +4*$B$6 includes two LayerNorms. Hover over those cells "
        "for the exact association. Activations, residual additions, dropout, "
        "softmax, and FFT arithmetic are omitted under the LRAC sheet convention. Overall "
        "Stride in B5 drives compute, field-of-view, and memory-bandwidth "
        "formulas. Channel dimensions are centralized in B6:B9, and bounded "
        "attention and RoPE settings are centralized in B10:B14. Bytes per value "
        "and runtime state read/write assumptions are centralized in B16:B18. The "
        "layer table references those cells. All totals derive only from visible "
        "formulas.",
    )
    sheet.cell(2, 4).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(2, 4).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    headers = [
        "Layer Name",
        "Block ID",
        "# Blocks",
        "Operation",
        "Kernel / Context",
        "Dilation",
        "Input Channels",
        "Output Channels",
        "Groups / Heads",
        "MAC / Position / Block",
        "MMAC/s",
        "Trainable Params / Block",
        "Total Trainable Params",
        "FOV Left / Block",
        "FOV Right / Block",
        "Total FOV Left (ms)",
        "Algorithmic Latency (ms)",
        "Parameter-Cap Values / Block",
        "Runtime State/Cache Accesses (R+W) / Call / Block",
        "Model-Params BW Total (MB/s)",
        "State/Cache BW Total (MB/s)",
        "Total BW (MB/s)",
    ]
    sheet.append(headers)
    header_row = sheet.max_row
    style_header(sheet[header_row])
    sheet.row_dimensions[header_row].height = 58

    first_component = header_row + 1
    previous_block = None
    shade_block = False

    def channel_cell(channels):
        channel_settings = {
            variant.model_dim: "=$B$6",
            (
                variant.encoder_convnext_expansion_dim
                if is_encoder
                else variant.decoder_convnext_expansion_dim
            ): "=$B$7",
            3 * variant.model_dim: "=$B$8",
            2 * (FFT_SIZE // 2 + 1): "=$B$9",
        }
        return channel_settings.get(channels, channels)

    def parameter_extra_cell(extra):
        if extra == 0:
            return "0"
        if extra % variant.model_dim == 0:
            return f"{extra // variant.model_dim}*$B$6"
        return str(extra)

    for layer in components:
        row = sheet.max_row + 1
        if layer.block != previous_block:
            shade_block = not shade_block
            previous_block = layer.block

        if layer.mac_kind == "conv":
            mac_cell = f"=E{row}*G{row}*H{row}/I{row}"
        elif layer.mac_kind == "attention":
            mac_cell = f"=I{row}*E{row}*(G{row}/I{row})"
        elif layer.mac_kind == "rope_phase":
            mac_cell = f"=E{row}"
        elif layer.mac_kind == "rope_rotation":
            mac_cell = f"=2*E{row}*I{row}"
        else:
            mac_cell = 0

        if layer.parameter_kind == "conv":
            parameter_cell = (
                f"=J{row}+H{row}+{parameter_extra_cell(layer.parameter_extra)}"
            )
        elif layer.parameter_kind == "norm":
            parameter_cell = f"=2*H{row}+{parameter_extra_cell(layer.parameter_extra)}"
        else:
            parameter_cell = layer.parameter_extra

        if layer.persistent_override is not None:
            persistent_cell = (
                "=$B$3"
                if layer.state_kind == "stft"
                else layer.persistent_override + layer.persistent_extra
            )
        elif layer.parameter_kind != "none":
            persistent_cell = f"=L{row}"
            if layer.persistent_extra:
                persistent_cell += f"+{layer.persistent_extra}"
        else:
            persistent_cell = layer.persistent_extra

        if layer.state_kind == "causal_conv":
            state_size = f"(E{row}-1)*F{row}*G{row}"
            state_cell = f"=$B$17*{state_size}+$B$18*{state_size}"
        elif layer.state_kind in ("attention_k", "attention_v"):
            state_cell = f"=$B$17*($B$10-1+$B$11)*G{row}+$B$18*G{row}"
        elif layer.state_kind == "stft":
            state_cell = "=$B$17*($B$3-$B$4)+$B$18*($B$3-$B$4)"
        elif layer.state_kind == "rope_inv_freq":
            state_cell = "=$B$17*($B$14/2)"
        else:
            state_cell = 0

        groups_or_heads = layer.groups_or_heads
        if layer.mac_kind in ("attention", "rope_rotation"):
            groups_or_heads = "=$B$12"
        elif groups_or_heads == variant.model_dim:
            groups_or_heads = "=$B$6"

        kernel_or_context = layer.kernel_or_context
        left_context = layer.left_context
        right_context = layer.right_context
        if layer.mac_kind == "attention":
            kernel_or_context = "=$B$10+$B$11"
        elif layer.mac_kind == "rope_phase":
            kernel_or_context = "=$B$14/2"
        elif layer.mac_kind == "rope_rotation":
            kernel_or_context = "=$B$14"
        elif layer.state_kind == "stft":
            kernel_or_context = "=$B$3"
        if layer.name == "Query-key score product":
            left_context = "=$B$10-1"
            right_context = "=$B$11"

        sheet.append(
            [
                layer.name,
                layer.block,
                layer.instances,
                layer.operation,
                kernel_or_context,
                layer.dilation,
                channel_cell(layer.input_channels),
                channel_cell(layer.output_channels),
                groups_or_heads,
                mac_cell,
                f"=C{row}*$B$5*J{row}/10^6",
                parameter_cell,
                f"=C{row}*L{row}",
                left_context,
                right_context,
                f"=C{row}*N{row}/$B$5*1000",
                f"=C{row}*O{row}/$B$5*1000",
                persistent_cell,
                state_cell,
                f"=C{row}*$B$5*$B$16*R{row}/10^6",
                f"=C{row}*$B$5*$B$16*S{row}/10^6",
                f"=SUM(T{row}:U{row})",
            ]
        )
        if layer.parameter_extra_note:
            sheet.cell(row, 12).comment = Comment(
                f"Includes {layer.parameter_extra_note}. The additive term in "
                "this formula is 2 × model channels per LayerNorm.",
                "LRAC 2",
            )
        if layer.state_kind in ("attention_k", "attention_v"):
            cache_name = "key" if layer.state_kind == "attention_k" else "value"
            sheet.cell(row, 19).comment = Comment(
                f"The first term reads the {cache_name} cache for the previous/"
                f"lookahead frames; the second writes the current {cache_name}. "
                "Attention context includes the current frame.",
                "LRAC 2",
            )
        elif layer.state_kind == "rope_inv_freq":
            sheet.cell(row, 19).comment = Comment(
                "The inverse-frequency vector is derived rather than learned and "
                "does not count toward the parameter cap. It is read once per "
                "attention invocation; the generated sine/cosine values are "
                "retained and reused for both query and key rotation.",
                "LRAC 2",
            )
        elif layer.state_kind in ("causal_conv", "stft"):
            sheet.cell(row, 19).comment = Comment(
                "The first term reads the runtime state and the second term "
                "writes it. The access counts come from cells B17 and B18.",
                "LRAC 2",
            )
        if shade_block:
            for cell in sheet[row]:
                cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        if layer.instances > 1:
            sheet.row_dimensions[row].outlineLevel = 1
        for column in (10, 12, 13, 18, 19):
            sheet.cell(row, column).number_format = "#,##0"
        for column in (11, 16, 17, 20, 21, 22):
            sheet.cell(row, column).number_format = "0.0000"

    last_component = sheet.max_row
    total_row = last_component + 1
    sheet.cell(total_row, 1, "Formula-based total")
    sheet.cell(total_row, 11, f"=SUM(K{first_component}:K{last_component})")
    sheet.cell(total_row, 13, f"=SUM(M{first_component}:M{last_component})")
    sheet.cell(
        total_row,
        16,
        f"=SUM(P{first_component}:P{last_component})",
    )
    sheet.cell(
        total_row,
        17,
        f"=SUM(Q{first_component}:Q{last_component})"
        + ("" if is_encoder else "+$B$15"),
    )
    sheet.cell(
        total_row,
        18,
        f"=SUMPRODUCT(C{first_component}:C{last_component},"
        f"R{first_component}:R{last_component})",
    )
    sheet.cell(
        total_row,
        20,
        f"=SUM(T{first_component}:T{last_component})",
    )
    sheet.cell(
        total_row,
        21,
        f"=SUM(U{first_component}:U{last_component})",
    )
    sheet.cell(total_row, 22, f"=SUM(T{total_row}:U{total_row})")
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=GREEN)

    mflops_row = total_row + 1
    sheet.cell(mflops_row, 1, "Total compute (MFLOPS)")
    sheet.cell(mflops_row, 11, f"=K{total_row}*2")
    for cell in sheet[mflops_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=BLUE)

    note_row = mflops_row + 2
    sheet.cell(
        note_row,
        1,
        "Memory bandwidth uses the bytes-per-value setting in B16 and includes "
        "model parameters, model buffers, convolution state, and attention "
        "KV-cache. Parameter-cap "
        "values include trainable parameters and non-runtime model buffers "
        "required for inference. Runtime state/cache is excluded from the "
        "parameter cap, but its read/write traffic is included in memory "
        "bandwidth. Intermediate activations are excluded by the challenge "
        "convention. Attention MACs "
        "assume that the configured context buffer is fully populated; startup "
        "frames require less compute. FOV left is history, not latency; only FOV "
        "right contributes to algorithmic latency. STFT/ISTFT window and "
        "overlap-state bandwidth is shown explicitly in the first table row; FFT "
        "arithmetic is excluded.",
    )
    sheet.merge_cells(
        start_row=note_row, start_column=1, end_row=note_row + 1, end_column=22
    )
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(note_row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    style_table(sheet, header_row, mflops_row, 1, 22)
    set_widths(
        sheet,
        [
            31,
            32,
            10,
            38,
            17,
            10,
            14,
            15,
            15,
            22,
            13,
            22,
            21,
            16,
            17,
            19,
            21,
            22,
            28,
            22,
            26,
            19,
        ],
    )
    sheet.freeze_panes = f"A{first_component}"
    sheet.auto_filter.ref = f"A{header_row}:V{last_component}"
    sheet.sheet_view.zoomScale = 70
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return total_row


def validate_supported_bitrates(variant):
    if not variant.supported_bitrates_kbps:
        raise ValueError("supported_bitrates_kbps must not be empty")
    bits_per_layer = math.log2(CODEBOOK_SIZE)
    for bitrate in variant.supported_bitrates_kbps:
        active_layers = bitrate * 1000 / (variant.frame_rate * bits_per_layer)
        if not active_layers.is_integer():
            raise ValueError(f"{bitrate} kbps is not representable by {variant.label}")
        if not 1 <= active_layers <= variant.rvq_layers:
            raise ValueError(
                f"{bitrate} kbps requires {int(active_layers)} RVQ layers, "
                f"outside the configured range 1..{variant.rvq_layers}"
            )


def validate_variant(variant):
    """Reject architecture settings that cannot describe the implementation."""

    if variant.model_dim % variant.attention_heads:
        raise ValueError("model_dim must be divisible by attention_heads")
    if variant.attention_use_rope:
        head_dim = variant.model_dim // variant.attention_heads
        rotary_dim = variant.attention_rotary_dim
        if rotary_dim is None or rotary_dim <= 0 or rotary_dim % 2:
            raise ValueError("attention_rotary_dim must be a positive even integer")
        if rotary_dim > head_dim:
            raise ValueError("attention_rotary_dim cannot exceed dimension per head")
    elif variant.attention_rotary_dim not in (None, 0):
        raise ValueError("attention_rotary_dim must be unset when RoPE is disabled")
    validate_supported_bitrates(variant)


def validate_config_matches_variant(variant, config_path):
    """Verify that a final recipe config and its design sheet describe one model."""

    import yaml

    with Path(config_path).open(encoding="utf-8") as stream:
        config = yaml.safe_load(stream)
    codec_conf = config["codec_conf"]
    generator = codec_conf["generator_params"]
    expected_decoder_groups = (tuple(variant.decoder_posnet_groups_first),) * 2 + (
        tuple(variant.decoder_posnet_groups_last),
    ) * 2
    expected = {
        "sampling_rate": (codec_conf["sampling_rate"], SAMPLE_RATE),
        "n_fft": (generator["n_fft"], FFT_SIZE),
        "hop_length": (generator["hop_length"], variant.hop_size),
        "hidden_dim": (generator["hidden_dim"], variant.model_dim),
        "encoder_num_convnext_blocks": (
            generator["encoder_num_convnext_blocks"],
            6,
        ),
        "encoder_convnext_hidden_dim": (
            generator["encoder_convnext_hidden_dim"],
            variant.encoder_convnext_expansion_dim,
        ),
        "decoder_num_convnext_blocks": (
            generator["decoder_num_convnext_blocks"],
            6,
        ),
        "decoder_convnext_hidden_dim": (
            generator["decoder_convnext_hidden_dim"],
            variant.decoder_convnext_expansion_dim,
        ),
        "convnext_kernel_size": (generator["convnext_kernel_size"], 7),
        "posnet_num_blocks": (generator["posnet_num_blocks"], 4),
        "posnet_kernel_size": (generator["posnet_kernel_size"], 3),
        "encoder_posnet_groups": (
            tuple(generator["encoder_posnet_groups"]),
            (1, 1),
        ),
        "decoder_posnet_groups": (
            tuple(tuple(groups) for groups in generator["decoder_posnet_groups"]),
            expected_decoder_groups,
        ),
        "attention_after": (generator["attention_after"], 2),
        "attention_context_frames": (
            generator["attention_context_frames"],
            variant.attention_context_frames,
        ),
        "attention_lookahead_frames": (
            generator["attention_lookahead_frames"],
            variant.attention_lookahead_frames,
        ),
        "attention_heads": (generator["attention_heads"], variant.attention_heads),
        "attention_use_rope": (
            generator.get("attention_use_rope", False),
            variant.attention_use_rope,
        ),
        "attention_rotary_dim": (
            generator.get("attention_rotary_dim"),
            variant.attention_rotary_dim,
        ),
        "attention_rope_base": (
            generator.get("attention_rope_base", 10000.0),
            10000.0,
        ),
        "latent_kernel_size": (generator["latent_kernel_size"], 3),
        "quantizer_codebook_dim": (generator["quantizer_codebook_dim"], CODEBOOK_DIM),
        "quantizer_n_q": (generator["quantizer_n_q"], variant.rvq_layers),
        "quantizer_bins": (generator["quantizer_bins"], CODEBOOK_SIZE),
        "quantizer_projection_mode": (
            generator["quantizer_projection_mode"],
            "per_layer",
        ),
        "quantizer_projection_bias": (
            generator["quantizer_projection_bias"],
            True,
        ),
        "quantizer_projection_norm": (
            generator["quantizer_projection_norm"],
            "weight_norm",
        ),
        "quantizer_target_bandwidth": (
            tuple(generator["quantizer_target_bandwidth"]),
            tuple(variant.supported_bitrates_kbps),
        ),
    }
    mismatches = [
        f"{name}: config={actual!r}, design={designed!r}"
        for name, (actual, designed) in expected.items()
        if actual != designed
    ]
    if mismatches:
        raise ValueError("Config/design mismatch:\n" + "\n".join(mismatches))


def add_vq_deployment_note(sheet, variant, last_column):
    maximum_bitrate = max(variant.supported_bitrates_kbps)
    rate_ceiling_note = ""
    if maximum_bitrate < max(SUPPORTED_BITRATES_KBPS):
        rate_ceiling_note = (
            f" This design natively supports {', '.join(map(str, variant.supported_bitrates_kbps))} "
            f"kbps. A request above {maximum_bitrate} kbps, including 32 kbps, "
            f"is clamped to all {variant.rvq_layers} layers and produces the same "
            f"output as {maximum_bitrate} kbps."
        )
    sheet.cell(
        7,
        1,
        f"The Track {variant.track} baseline uses one factorized DAC RVQ: the "
        "low-dimensional codebooks and output projections are shared logical "
        "parameters and are counted once, while their transmit- and receive-side "
        "accesses are counted separately. An alternative deployment may keep the "
        "low-dimensional transmit codebooks and materialize separate receive "
        "codebooks containing the output-projected vectors. That alternative "
        "removes receive output-projection compute and parameter-read bandwidth. "
        "Under this sheet's MAC convention, receive VQ compute then becomes zero; "
        "receive VQ bandwidth becomes Frame Rate x Bytes per Value x RVQ Layers x "
        "Model Dimension, and the parameter cap gains RVQ Layers x Codebook Size x "
        "Model Dimension values for the materialized receive codebooks. Transmit "
        "accounting is unchanged. Those alternative values are not included here."
        + rate_ceiling_note,
    )
    sheet.merge_cells(start_row=7, start_column=1, end_row=10, end_column=last_column)
    sheet.cell(7, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(7, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)


def create_vq_transmit_sheet(workbook, variant):
    validate_supported_bitrates(variant)
    encoder_sheet = variant.sheet_name("Encoder")
    sheet = workbook.create_sheet(variant.sheet_name("VQ-Transmit"))
    style_title(sheet, f"VoCodec DAC-Style RVQ Transmit — {variant.label}", 21)
    sheet.append([])
    headers = [
        "Embedding Dimension",
        "Projection Dimension",
        "Codebook Size",
        "Layers",
        "Frame Rate",
        "MMAC/s Input Projections",
        "MMAC/s Distance",
        "MMAC/s Output Projections",
        "Total Compute (MFLOPS)",
        "Parameters (biases and weight norm excluded)",
        "Unique VQ Trainable Parameters",
        "Maximum Native Bitrate (kbps)",
        "Latency (ms)",
        "Input Projection Values / Layer",
        "Output Projection Values / Layer",
        "Codebook Values / Layer",
        "Input Projection BW / Layer (MB/s)",
        "Output Projection BW / Layer (MB/s)",
        "Codebook Scan BW / Layer (MB/s)",
        "Total BW / Layer (MB/s)",
        "All Active Layers BW (MB/s)",
    ]
    sheet.append(headers)
    style_header(sheet[3])
    sheet.append(
        [
            f"='{encoder_sheet}'!$B$6",
            CODEBOOK_DIM,
            CODEBOOK_SIZE,
            variant.rvq_layers,
            f"='{encoder_sheet}'!$B$5",
            "=A4*B4*D4*E4/10^6",
            "=B4*C4*D4*E4/10^6",
            "=B4*A4*D4*E4/10^6",
            "=SUM(F4:H4)*2",
            "=D4*(C4*B4+A4*B4+B4*A4)",
            "=D4*(N4+O4+P4)",
            "=LOG(C4,2)*D4*E4/1000",
            0,
            "=A4*B4+2*B4",
            "=B4*A4+2*A4",
            "=B4*C4",
            f"=E4*'{encoder_sheet}'!$B$16*N4/10^6",
            f"=E4*'{encoder_sheet}'!$B$16*O4/10^6",
            f"=E4*'{encoder_sheet}'!$B$16*P4/10^6",
            "=SUM(Q4:S4)",
            "=T4*D4",
        ]
    )
    style_table(sheet, 3, 4, 1, 21)
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    add_vq_deployment_note(sheet, variant, 21)

    set_widths(
        sheet,
        [
            19,
            21,
            16,
            10,
            13,
            21,
            18,
            22,
            25,
            25,
            26,
            19,
            23,
            13,
            23,
            24,
            21,
            25,
            24,
            23,
            27,
        ],
    )
    sheet.freeze_panes = "A4"


def create_vq_receive_sheet(workbook, variant):
    transmit_sheet = variant.sheet_name("VQ-Transmit")
    encoder_sheet = variant.sheet_name("Encoder")
    sheet = workbook.create_sheet(variant.sheet_name("VQ-Receive"))
    style_title(sheet, f"VoCodec DAC-Style RVQ Receive — {variant.label}", 16)
    sheet.append([])
    headers = [
        "Embedding Dimension",
        "Projection Dimension",
        "Codebook Size",
        "Active Layers at Maximum Bitrate",
        "Frame Rate",
        "MMAC/s Output Projections",
        "Total Compute (MFLOPS)",
        "Shared VQ Parameters (counted on transmit tab)",
        "Additional Unique Parameters",
        "Latency (ms)",
        "Selected Embedding Values / Layer",
        "Output Projection Values / Layer",
        "Selected Embedding BW / Layer (MB/s)",
        "Output Projection BW / Layer (MB/s)",
        "Total BW / Layer (MB/s)",
        "All Active Layers BW (MB/s)",
    ]
    sheet.append(headers)
    style_header(sheet[3])
    sheet.append(
        [
            f"='{transmit_sheet}'!A4",
            f"='{transmit_sheet}'!B4",
            f"='{transmit_sheet}'!C4",
            f"='{transmit_sheet}'!D4",
            f"='{transmit_sheet}'!E4",
            "=A4*B4*D4*E4/10^6",
            "=F4*2",
            f"='{transmit_sheet}'!K4",
            0,
            0,
            "=B4",
            f"='{transmit_sheet}'!O4",
            f"=E4*'{encoder_sheet}'!$B$16*K4/10^6",
            f"=E4*'{encoder_sheet}'!$B$16*L4/10^6",
            "=SUM(M4:N4)",
            "=O4*D4",
        ]
    )
    style_table(sheet, 3, 4, 1, 16)
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    add_vq_deployment_note(sheet, variant, 16)
    set_widths(
        sheet,
        [
            19,
            21,
            16,
            25,
            13,
            26,
            23,
            29,
            24,
            14,
            26,
            27,
            28,
            29,
            23,
            27,
        ],
    )
    sheet.freeze_panes = "A4"


def bandwidth_penalty_formula(bandwidth_cell, pmax_cell):
    return (
        f"=IF({bandwidth_cell}<=50,0,"
        f"IF({bandwidth_cell}>=1000,{pmax_cell},"
        f"{pmax_cell}*LN({bandwidth_cell}/50)/LN(1000/50)))"
    )


def create_summary_sheet(workbook, variant, encoder_total_row, decoder_total_row):
    encoder_sheet = variant.sheet_name("Encoder")
    vq_transmit_sheet = variant.sheet_name("VQ-Transmit")
    vq_receive_sheet = variant.sheet_name("VQ-Receive")
    decoder_sheet = variant.sheet_name("Decoder")
    sheet = workbook.create_sheet(variant.sheet_name("Summary"))
    style_title(
        sheet,
        f"VoCodec LRAC 2 Track {variant.track} Summary — {variant.label}",
        5,
    )
    sheet.append([])
    headers = ["Measurement", "Model", "Budget", "Margin", "Status"]
    sheet.append(headers)
    style_header(sheet[3])
    rows = [
        [
            "Transmit compute (MFLOPS)",
            f"='{encoder_sheet}'!K{encoder_total_row}*2" f"+'{vq_transmit_sheet}'!I4",
            TRANSMIT_BUDGET_MFLOPS,
            "=C4-B4",
            '=IF(B4<=C4,"PASS","OVER BUDGET")',
        ],
        [
            "Receive compute (MFLOPS)",
            f"='{vq_receive_sheet}'!G4" f"+'{decoder_sheet}'!K{decoder_total_row}*2",
            RECEIVE_BUDGET_MFLOPS,
            "=C5-B5",
            '=IF(B5<=C5,"PASS","OVER BUDGET")',
        ],
        [
            "Buffering latency (ms)",
            f"='{encoder_sheet}'!$B$15",
            "",
            "",
            "",
        ],
        [
            "Algorithmic latency (ms)",
            f"='{encoder_sheet}'!Q{encoder_total_row}"
            f"+'{decoder_sheet}'!Q{decoder_total_row}",
            "",
            "",
            "",
        ],
        [
            "Total latency (ms)",
            "=SUM(B6:B7)",
            LATENCY_BUDGET_MS,
            "=C8-B8",
            '=IF(B8<=C8,"PASS","OVER BUDGET")',
        ],
        [
            "Maximum native bitrate (kbps)",
            f"='{vq_transmit_sheet}'!L4",
            "",
            "",
            "",
        ],
        [
            "Generator trainable parameters",
            f"='{encoder_sheet}'!M{encoder_total_row}"
            f"+'{vq_transmit_sheet}'!K4"
            f"+'{decoder_sheet}'!M{decoder_total_row}",
            "",
            "",
            "",
        ],
        [
            "Parameter-cap values (parameters + model buffers)",
            f"='{encoder_sheet}'!R{encoder_total_row}"
            f"+'{vq_transmit_sheet}'!K4"
            f"+'{decoder_sheet}'!R{decoder_total_row}",
            variant.parameter_budget or "",
            "=C11-B11" if variant.parameter_budget else "",
            '=IF(B11<=C11,"PASS","OVER BUDGET")' if variant.parameter_budget else "",
        ],
    ]
    for row in rows:
        sheet.append(row)
    style_table(sheet, 3, 11, 1, 5)
    for row in (4, 5, 8):
        sheet.cell(row, 5).fill = PatternFill("solid", fgColor=GREEN)
        sheet.cell(row, 5).font = Font(bold=True)
    if variant.parameter_budget:
        sheet.cell(11, 5).fill = PatternFill("solid", fgColor=GREEN)
        sheet.cell(11, 5).font = Font(bold=True)

    memory_start = 13
    sheet.cell(memory_start, 1, "MEMORY BANDWIDTH AT MAXIMUM BITRATE")
    sheet.merge_cells(
        start_row=memory_start,
        start_column=1,
        end_row=memory_start,
        end_column=5,
    )
    sheet.cell(memory_start, 1).font = Font(bold=True, color=WHITE)
    sheet.cell(memory_start, 1).fill = PatternFill("solid", fgColor=NAVY)
    memory_rows = [
        ["Encoder memory bandwidth (MB/s)", f"='{encoder_sheet}'!V{encoder_total_row}"],
        ["RVQ transmit memory bandwidth (MB/s)", f"='{vq_transmit_sheet}'!U4"],
        ["RVQ receive memory bandwidth (MB/s)", f"='{vq_receive_sheet}'!P4"],
        ["Decoder memory bandwidth (MB/s)", f"='{decoder_sheet}'!V{decoder_total_row}"],
        ["Total codec memory bandwidth (MB/s)", "=SUM(B14:B17)"],
        ["Penalty maximum P_max", MEMORY_BANDWIDTH_PMAX],
        ["Bandwidth penalty P(B)", bandwidth_penalty_formula("B18", "$B$19")],
    ]
    for row in memory_rows:
        sheet.append(row)
    style_table(sheet, 14, 20, 1, 5)
    sheet["B19"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    sheet["B19"].font = Font(bold=True)
    for row in range(14, 21):
        sheet[f"B{row}"].number_format = "0.0000"

    conclusion_row = 23
    sheet.cell(
        conclusion_row,
        1,
        f"The {variant.label} design meets the transmitter compute, receiver "
        f"compute, and total-latency Track {variant.track} limits. "
        + (
            "It also meets the parameter-cap value limit. "
            if variant.parameter_budget
            else ""
        )
        + "Memory bandwidth includes encoder, maximum-native-rate RVQ transmit "
        "and receive paths, and decoder."
        + (
            f" Requests above {max(variant.supported_bitrates_kbps)} kbps are "
            "clamped to the maximum native RVQ prefix."
            if max(variant.supported_bitrates_kbps) < max(SUPPORTED_BITRATES_KBPS)
            else ""
        ),
    )
    sheet.merge_cells(
        start_row=conclusion_row,
        start_column=1,
        end_row=conclusion_row + 1,
        end_column=5,
    )
    sheet.cell(conclusion_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.cell(conclusion_row, 1).fill = PatternFill("solid", fgColor=GREEN)
    set_widths(sheet, [45, 26, 18, 18, 20])
    sheet.freeze_panes = "A4"


def calculate_variant(variant):
    input_projection_values = variant.model_dim * CODEBOOK_DIM + 2 * CODEBOOK_DIM
    output_projection_values = CODEBOOK_DIM * variant.model_dim + 2 * variant.model_dim
    codebook_values = CODEBOOK_SIZE * CODEBOOK_DIM
    rvq_parameters_per_layer = (
        input_projection_values + output_projection_values + codebook_values
    )
    input_projection = (
        variant.frame_rate * variant.model_dim * CODEBOOK_DIM * variant.rvq_layers / 1e6
    )
    distance = (
        variant.frame_rate * CODEBOOK_DIM * CODEBOOK_SIZE * variant.rvq_layers / 1e6
    )
    vq_transmit_mflops = 2 * (2 * input_projection + distance)
    vq_receive_mflops = 2 * input_projection

    def layer_mmac(components):
        return sum(
            layer.instances * variant.frame_rate * layer.macs_per_position() / 1e6
            for layer in components
        )

    def memory_bandwidth(components):
        return sum(
            layer.instances
            * variant.frame_rate
            * BYTES_PER_VALUE
            * (layer.persistent_values(variant) + layer.state_values(variant))
            / 1e6
            for layer in components
        )

    def algorithmic_latency_ms(components):
        return sum(
            layer.instances
            * layer.right_context
            / (SAMPLE_RATE if layer.context_domain == "sample" else variant.frame_rate)
            * 1000
            for layer in components
        )

    encoder_layers = encoder_components(variant)
    decoder_layers = decoder_components(variant)
    encoder_parameters = sum(
        layer.instances * layer.trainable_parameters() for layer in encoder_layers
    )
    decoder_parameters = sum(
        layer.instances * layer.trainable_parameters() for layer in decoder_layers
    )
    encoder_persistent_values = sum(
        layer.instances * layer.persistent_values(variant) for layer in encoder_layers
    )
    decoder_persistent_values = sum(
        layer.instances * layer.persistent_values(variant) for layer in decoder_layers
    )
    encoder_layer_mmac = layer_mmac(encoder_layers)
    decoder_layer_mmac = layer_mmac(decoder_layers)
    encoder_bandwidth = memory_bandwidth(encoder_layers)
    decoder_bandwidth = memory_bandwidth(decoder_layers)
    rvq_transmit_bandwidth = (
        variant.frame_rate
        * BYTES_PER_VALUE
        * rvq_parameters_per_layer
        * variant.rvq_layers
        / 1e6
    )
    rvq_receive_bandwidth = (
        variant.frame_rate
        * BYTES_PER_VALUE
        * (output_projection_values + CODEBOOK_DIM)
        * variant.rvq_layers
        / 1e6
    )
    total_bandwidth = (
        encoder_bandwidth
        + rvq_transmit_bandwidth
        + rvq_receive_bandwidth
        + decoder_bandwidth
    )
    penalty = (
        0.0
        if total_bandwidth <= 50
        else (
            float(MEMORY_BANDWIDTH_PMAX)
            if total_bandwidth >= 1000
            else MEMORY_BANDWIDTH_PMAX
            * math.log(total_bandwidth / 50)
            / math.log(1000 / 50)
        )
    )
    encoder_latency_ms = algorithmic_latency_ms(encoder_layers)
    decoder_latency_ms = algorithmic_latency_ms(decoder_layers) + (
        (FFT_SIZE - variant.hop_size) / SAMPLE_RATE * 1000
    )
    return {
        "encoder_formula_mmac": encoder_layer_mmac,
        "encoder_mflops": 2 * encoder_layer_mmac,
        "vq_transmit_mflops": vq_transmit_mflops,
        "vq_receive_mflops": vq_receive_mflops,
        "transmit_mflops": 2 * encoder_layer_mmac + vq_transmit_mflops,
        "decoder_formula_mmac": decoder_layer_mmac,
        "receive_mflops": vq_receive_mflops + 2 * decoder_layer_mmac,
        "buffering_latency_ms": variant.hop_size / SAMPLE_RATE * 1000,
        "encoder_algorithmic_latency_ms": encoder_latency_ms,
        "decoder_algorithmic_latency_ms": decoder_latency_ms,
        "latency_ms": (
            variant.hop_size / SAMPLE_RATE * 1000
            + encoder_latency_ms
            + decoder_latency_ms
        ),
        "maximum_bitrate_kbps": (
            variant.rvq_layers * variant.frame_rate * math.log2(CODEBOOK_SIZE) / 1000
        ),
        "generator_parameters": (
            encoder_parameters
            + decoder_parameters
            + rvq_parameters_per_layer * variant.rvq_layers
        ),
        "deployed_persistent_values": (
            encoder_persistent_values
            + decoder_persistent_values
            + rvq_parameters_per_layer * variant.rvq_layers
        ),
        "encoder_memory_bandwidth_mbps": encoder_bandwidth,
        "rvq_transmit_memory_bandwidth_mbps": rvq_transmit_bandwidth,
        "rvq_receive_memory_bandwidth_mbps": rvq_receive_bandwidth,
        "rvq_memory_bandwidth_mbps": (rvq_transmit_bandwidth + rvq_receive_bandwidth),
        "decoder_memory_bandwidth_mbps": decoder_bandwidth,
        "total_memory_bandwidth_mbps": total_bandwidth,
        "bandwidth_penalty": penalty,
    }


def validate_persistent_value_formulas(
    workbook, variant, encoder_total_row, decoder_total_row
):
    """Check workbook persistent-value accounting against the Python model."""

    component_totals = {}
    for component_name, components, total_row in (
        ("Encoder", encoder_components(variant), encoder_total_row),
        ("Decoder", decoder_components(variant), decoder_total_row),
    ):
        sheet = workbook[variant.sheet_name(component_name)]
        first_row = total_row - len(components)
        last_row = total_row - 1
        expected_total_formula = (
            f"=SUMPRODUCT(C{first_row}:C{last_row},R{first_row}:R{last_row})"
        )
        if sheet.cell(total_row, 18).value != expected_total_formula:
            raise ValueError(
                f"{component_name} persistent total must include block multiplicity"
            )

        component_total = 0
        for row, layer in zip(range(first_row, total_row), components):
            if sheet.cell(row, 3).value != layer.instances:
                raise ValueError(f"{component_name} block count mismatch at row {row}")

            if layer.persistent_override is not None:
                expected_cell = (
                    "=$B$3"
                    if layer.state_kind == "stft"
                    else layer.persistent_override + layer.persistent_extra
                )
            elif layer.parameter_kind != "none":
                expected_cell = f"=L{row}"
                if layer.persistent_extra:
                    expected_cell += f"+{layer.persistent_extra}"
            else:
                expected_cell = layer.persistent_extra
            if sheet.cell(row, 18).value != expected_cell:
                raise ValueError(
                    f"{component_name} persistent-value formula mismatch at row {row}"
                )

            component_total += layer.instances * layer.persistent_values(variant)
        component_totals[component_name] = component_total

    summary = workbook[variant.sheet_name("Summary")]
    expected_summary_formula = (
        f"='{variant.sheet_name('Encoder')}'!R{encoder_total_row}"
        f"+'{variant.sheet_name('VQ-Transmit')}'!K4"
        f"+'{variant.sheet_name('Decoder')}'!R{decoder_total_row}"
    )
    if summary["B11"].value != expected_summary_formula:
        raise ValueError("Summary must use the component persistent-value totals")

    rvq_total = (
        CODEBOOK_SIZE * CODEBOOK_DIM
        + 2 * variant.model_dim * CODEBOOK_DIM
        + 2 * CODEBOOK_DIM
        + 2 * variant.model_dim
    ) * variant.rvq_layers
    workbook_total = (
        component_totals["Encoder"] + rvq_total + component_totals["Decoder"]
    )
    python_total = calculate_variant(variant)["deployed_persistent_values"]
    if workbook_total != python_total:
        raise ValueError(
            "Workbook and Python deployed persistent-value calculations disagree"
        )


def validate_vq_formulas(workbook, variant, encoder_total_row, decoder_total_row):
    """Check that both factorized VQ paths feed the summary exactly once."""

    transmit_name = variant.sheet_name("VQ-Transmit")
    receive_name = variant.sheet_name("VQ-Receive")
    encoder_name = variant.sheet_name("Encoder")
    decoder_name = variant.sheet_name("Decoder")
    transmit = workbook[transmit_name]
    receive = workbook[receive_name]
    summary = workbook[variant.sheet_name("Summary")]

    for vq_sheet in (transmit, receive):
        if any(cell.value == "Calls/s" for row in vq_sheet.iter_rows() for cell in row):
            raise ValueError("VQ sheets must use the frame-rate convention")

    expected_cells = {
        (transmit, "I4"): "=SUM(F4:H4)*2",
        (transmit, "K4"): "=D4*(N4+O4+P4)",
        (transmit, "U4"): "=T4*D4",
        (receive, "G4"): "=F4*2",
        (receive, "H4"): f"='{transmit_name}'!K4",
        (receive, "I4"): 0,
        (receive, "P4"): "=O4*D4",
        (summary, "B4"): (
            f"='{encoder_name}'!K{encoder_total_row}*2+'{transmit_name}'!I4"
        ),
        (summary, "B5"): (
            f"='{receive_name}'!G4+'{decoder_name}'!K{decoder_total_row}*2"
        ),
        (summary, "B15"): f"='{transmit_name}'!U4",
        (summary, "B16"): f"='{receive_name}'!P4",
        (summary, "B18"): "=SUM(B14:B17)",
    }
    for (sheet, coordinate), expected in expected_cells.items():
        if sheet[coordinate].value != expected:
            raise ValueError(
                f"Unexpected factorized VQ formula in {sheet.title}!{coordinate}"
            )


def create_workbook(variant, output):
    """Create one workbook at the explicitly requested output path."""

    output = Path(output)
    validate_variant(variant)
    workbook = Workbook()
    del workbook[workbook.sheetnames[0]]
    encoder_total_row = add_component_sheet(
        workbook,
        variant,
        variant.sheet_name("Encoder"),
        encoder_components(variant),
    )
    create_vq_transmit_sheet(workbook, variant)
    create_vq_receive_sheet(workbook, variant)
    decoder_total_row = add_component_sheet(
        workbook,
        variant,
        variant.sheet_name("Decoder"),
        decoder_components(variant),
    )
    create_summary_sheet(workbook, variant, encoder_total_row, decoder_total_row)
    validate_persistent_value_formulas(
        workbook, variant, encoder_total_row, decoder_total_row
    )
    validate_vq_formulas(workbook, variant, encoder_total_row, decoder_total_row)

    workbook.properties.title = (
        f"LRAC 2 Track {variant.track} VoCodec {variant.label} Design"
    )
    workbook.properties.subject = "Compute, latency, parameters, and memory bandwidth"
    workbook.properties.description = (
        "VoCodec calculations using the LRAC compute and persistent-tensor "
        "memory-bandwidth conventions."
    )
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    results = calculate_variant(variant)
    assert results["receive_mflops"] <= RECEIVE_BUDGET_MFLOPS
    assert results["transmit_mflops"] <= TRANSMIT_BUDGET_MFLOPS
    assert results["latency_ms"] <= LATENCY_BUDGET_MS
    assert math.isclose(
        results["maximum_bitrate_kbps"], max(variant.supported_bitrates_kbps)
    )
    if variant.parameter_budget is not None:
        assert results["deployed_persistent_values"] <= variant.parameter_budget
    return results
